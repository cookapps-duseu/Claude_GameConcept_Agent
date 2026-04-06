"""
게임 컨셉 파이프라인 출력 스크립트.

Usage:
    python scripts/export.py --session output/session.json
    python scripts/export.py --session output/session.json --word
    python scripts/export.py --session output/session.json --template-dir templates
"""
import argparse
import html as html_mod
import json
import re
from datetime import datetime
from pathlib import Path

from jinja2 import Environment, FileSystemLoader


# ---------------------------------------------------------------------------
# 씬 파싱
# ---------------------------------------------------------------------------

def parse_concept_scenes(text: str) -> tuple[str, list[dict]]:
    """컨셉 텍스트에서 ---SCENE_LIST--- 이후 JSON을 파싱해 씬 목록을 반환.

    Returns:
        (clean_concept_text, scenes_list)
        씬 파싱 실패 시 (---SCENE_LIST--- 이전 텍스트, []) 반환.
    """
    delimiter = "---SCENE_LIST---"
    if delimiter not in text:
        return text, []
    parts = text.split(delimiter, 1)
    clean_text = parts[0].rstrip()
    scene_json = parts[1].strip()
    scene_json = re.sub(r"^```(?:json)?\s*", "", scene_json)
    scene_json = re.sub(r"\s*```$", "", scene_json)
    try:
        scenes = json.loads(scene_json)
        return clean_text, scenes if isinstance(scenes, list) else []
    except json.JSONDecodeError:
        return clean_text, []


# ---------------------------------------------------------------------------
# 게임 타이틀 추출
# ---------------------------------------------------------------------------

def extract_game_title(concept: str, genre: str) -> str:
    """컨셉 텍스트에서 게임 타이틀을 추출해 폴더명용 문자열로 반환.

    '게임 타이틀 (가제):' 패턴 → 괄호 제거 → 특수문자 제거 → 최대 30자.
    추출 실패 시 장르명 사용.
    """
    match = re.search(r'게임\s*타이틀[^::\uff1a]*[::\uff1a]\s*(.+)', concept)
    if match:
        title = match.group(1).strip()
        title = re.sub(r'\*+', '', title)               # 마크다운 볼드 제거
        title = re.sub(r'\s*\([^)]*\)\s*$', '', title).strip()  # (가제) 등 제거
        title = re.sub(r'[\\/:*?"<>|]', '', title)      # 파일시스템 특수문자 제거
        title = title[:30].strip()
        if title:
            return title
    sanitized = re.sub(r'[\\/:*?"<>|]', '', genre)
    return sanitized[:30].strip() or "concept"


# ---------------------------------------------------------------------------
# 마크다운 → HTML 변환
# ---------------------------------------------------------------------------

def table_to_html(table_lines: list[str]) -> str:
    """마크다운 표를 HTML <table class="md-table">로 변환."""
    if not table_lines:
        return ""

    def parse_row(line: str) -> list[str]:
        return [c.strip() for c in line.strip().strip("|").split("|")]

    rows = [parse_row(line) for line in table_lines]
    has_header = (
        len(rows) >= 2
        and all(re.match(r"^:?-+:?$", c) for c in rows[1] if c)
    )

    parts = ['<table class="md-table">']
    if has_header:
        parts.append("<thead><tr>")
        for cell in rows[0]:
            parts.append(f"<th>{html_mod.escape(cell)}</th>")
        parts.append("</tr></thead><tbody>")
        data_rows = rows[2:]
    else:
        parts.append("<tbody>")
        data_rows = rows

    for row in data_rows:
        parts.append("<tr>")
        for cell in row:
            parts.append(f"<td>{html_mod.escape(cell)}</td>")
        parts.append("</tr>")
    parts.append("</tbody></table>")
    return "\n".join(parts)


def markdown_to_html(text: str) -> str:
    """마크다운 텍스트를 HTML로 변환. 표, 헤딩, 리스트 지원."""
    lines = text.split("\n")
    html_lines = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2:
            table_lines = []
            while i < len(lines):
                sl = lines[i].strip()
                if sl.startswith("|") and sl.endswith("|"):
                    table_lines.append(lines[i])
                    i += 1
                else:
                    break
            html_lines.append(table_to_html(table_lines))
            continue
        if line.startswith("# "):
            html_lines.append(f"<h1>{line[2:]}</h1>")
        elif line.startswith("## "):
            html_lines.append(f"<h2>{line[3:]}</h2>")
        elif line.startswith("### "):
            html_lines.append(f"<h3>{line[4:]}</h3>")
        elif line.startswith("- "):
            html_lines.append(f"<li>{line[2:]}</li>")
        elif line == "":
            html_lines.append("<br>")
        else:
            html_lines.append(f"<p>{line}</p>")
        i += 1
    return "\n".join(html_lines)


def inline_svgs(concept_html: str, svgs: list[dict]) -> str:
    """[IMG_SCENE_N] 플레이스홀더를 SVG div로 치환. 플레이스홀더 없는 씬은 끝에 추가."""
    if not svgs:
        return concept_html

    svg_map = {s["id"]: s for s in svgs}
    used_ids: set[int] = set()

    def replace_placeholder(m: re.Match) -> str:
        n = int(m.group(1))
        used_ids.add(n)
        scene = svg_map.get(n)
        if not scene or not scene.get("svg"):
            return ""
        return (
            f'<div class="ui-mockup">'
            f'<p class="scene-title">{html_mod.escape(scene.get("title", ""))}</p>'
            f'{scene["svg"]}'
            f'</div>'
        )

    result = re.sub(r'\[IMG_SCENE_(\d+)\]', replace_placeholder, concept_html)

    for scene in svgs:
        if scene["id"] not in used_ids and scene.get("svg"):
            result += (
                f'<div class="ui-mockup">'
                f'<p class="scene-title">{html_mod.escape(scene.get("title", ""))}</p>'
                f'{scene["svg"]}'
                f'</div>'
            )
    return result


def split_concept_sections(concept_html: str) -> dict[str, str]:
    """h2 태그 기준으로 섹션 분리. 키: "1"~"6". 실패 시 {"all": concept_html}."""
    pattern = re.compile(r'(<h2>[^<]*</h2>)', re.IGNORECASE)
    parts = pattern.split(concept_html)
    if len(parts) <= 1:
        return {"all": concept_html}

    sections: dict[str, str] = {}
    current_num = None
    buffer: list[str] = []

    for part in parts:
        h2_match = re.match(r'<h2>(\d+)\.\s', part, re.IGNORECASE)
        if h2_match:
            if current_num is not None:
                sections[current_num] = "".join(buffer)
            current_num = h2_match.group(1)
            buffer = [part]
        else:
            if current_num is None:
                sections["0"] = sections.get("0", "") + part
            else:
                buffer.append(part)

    if current_num is not None:
        sections[current_num] = "".join(buffer)

    return sections if sections else {"all": concept_html}


# ---------------------------------------------------------------------------
# 섹션 추출 (core_loop.md, core_fun.md용)
# ---------------------------------------------------------------------------

def extract_section(text: str, section_num: int) -> str:
    """마크다운 텍스트에서 '## N. ...' 섹션을 추출한다."""
    pattern = rf'## {section_num}\..+?(?=\n## \d+\.|\Z)'
    match = re.search(pattern, text, re.DOTALL)
    return match.group(0).strip() if match else ""


# ---------------------------------------------------------------------------
# 내보내기 함수
# ---------------------------------------------------------------------------

def export_html(session: dict, output_dir: Path, template_dir: str = "templates") -> Path:
    """concept.html 생성."""
    concept = session.get("concept", "")
    clean_concept, scenes = parse_concept_scenes(concept)
    concept_html = markdown_to_html(clean_concept)
    concept_html = inline_svgs(concept_html, session.get("concept_ui_svgs", []))

    user_input = session.get("user_input", {})

    env = Environment(loader=FileSystemLoader(template_dir))
    template = env.get_template("concept.html.j2")
    html = template.render(
        genre=user_input.get("genre", ""),
        platform=user_input.get("platform"),
        generated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        concept_html=concept_html,
        reference_images=session.get("reference_images", []),
        concept_diagram_mermaid=session.get("concept_diagram_mermaid", ""),
    )
    output_path = output_dir / "concept.html"
    output_path.write_text(html, encoding="utf-8")
    return output_path


def export_core_loop_md(session: dict, output_dir: Path) -> Path:
    """core_loop.md 생성 — 컨셉안 루프 섹션 + 레퍼런스 루프 분석."""
    loop_analysis = session.get("loop_analysis", "")
    concept_raw = session.get("concept", "").split("---SCENE_LIST---")[0]
    loop_section = extract_section(concept_raw, 4)

    lines = ["# 핵심 게임 루프\n"]
    if loop_section:
        lines.append(loop_section)
        lines.append("\n\n---\n")
    if loop_analysis:
        lines.append("\n## 레퍼런스 루프 분석\n")
        lines.append(loop_analysis)

    output_path = output_dir / "core_loop.md"
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def export_core_fun_md(session: dict, output_dir: Path) -> Path:
    """core_fun.md 생성 — 컨셉안 재미 섹션 + 레퍼런스 재미 분석."""
    fun_analysis = session.get("fun_analysis", "")
    concept_raw = session.get("concept", "").split("---SCENE_LIST---")[0]
    fun_section = extract_section(concept_raw, 3)

    lines = ["# 핵심 재미 요소\n"]
    if fun_section:
        lines.append(fun_section)
        lines.append("\n\n---\n")
    if fun_analysis:
        lines.append("\n## 레퍼런스 재미 분석\n")
        lines.append(fun_analysis)

    output_path = output_dir / "core_fun.md"
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def export_word(session: dict, output_dir: Path) -> Path:
    """concept.docx 생성."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    concept = session.get("concept", "").split("---SCENE_LIST---")[0]
    user_input = session.get("user_input", {})

    doc = Document()
    title_para = doc.add_heading("게임 컨셉안", level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "생성일: {} | 장르: {}".format(
            datetime.now().strftime("%Y-%m-%d"),
            user_input.get("genre", "")
        )
    )
    doc.add_paragraph()

    for line in concept.split("\n"):
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)

    output_path = output_dir / "concept.docx"
    doc.save(str(output_path))
    return output_path


# ---------------------------------------------------------------------------
# 메인
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description="게임 컨셉 파이프라인 출력 스크립트")
    parser.add_argument("--session", required=True, help="session.json 경로")
    parser.add_argument("--word", action="store_true", help="Word(.docx) 파일도 생성")
    parser.add_argument("--template-dir", default="templates", help="Jinja2 템플릿 디렉토리")
    args = parser.parse_args()

    session_path = Path(args.session)
    if not session_path.exists():
        raise FileNotFoundError(f"session.json을 찾을 수 없습니다: {session_path}")

    session = json.loads(session_path.read_text(encoding="utf-8"))
    user_input = session.get("user_input", {})
    genre = user_input.get("genre", "concept")
    concept = session.get("concept", "")

    title = extract_game_title(concept, genre)
    date_str = datetime.now().strftime("%Y%m%d")
    folder_name = f"{date_str}_{title}"

    output_dir = session_path.parent / folder_name
    output_dir.mkdir(exist_ok=True)

    export_html(session, output_dir, args.template_dir)
    export_core_loop_md(session, output_dir)
    export_core_fun_md(session, output_dir)

    if args.word:
        export_word(session, output_dir)

    # session.json을 출력 폴더로 이동
    dest = output_dir / "session.json"
    dest.write_text(session_path.read_text(encoding="utf-8"), encoding="utf-8")
    session_path.unlink()

    print(f"출력 완료: {output_dir}")


if __name__ == "__main__":
    main()
